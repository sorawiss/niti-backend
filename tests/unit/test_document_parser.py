from io import BytesIO
import unittest

from docx import Document
from docx.document import Document as DocxDocument

from app.documents.parser import parse_docx


def make_docx(document: DocxDocument) -> bytes:
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class DocumentParserTests(unittest.TestCase):
    def test_preserves_thai_paragraph_and_table_order(self) -> None:
        document = Document()
        document.add_paragraph("ข้อ 1 การจ้างงาน")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "เงินเดือน"
        table.cell(0, 1).text = "20,000 บาท"
        document.add_paragraph("ข้อ 2 เวลาทำงาน")

        blocks = parse_docx(make_docx(document))

        self.assertEqual(len(blocks), 3)
        self.assertEqual(
            [block.type for block in blocks],
            ["paragraph", "table", "paragraph"],
        )
        self.assertEqual(blocks[0].text, "ข้อ 1 การจ้างงาน")
        self.assertEqual(blocks[1].text, "เงินเดือน | 20,000 บาท")
        self.assertEqual(blocks[2].text, "ข้อ 2 เวลาทำงาน")
        self.assertEqual([block.index for block in blocks], [0, 1, 2])

    def test_classifies_heading_and_skips_empty_paragraphs(self) -> None:
        document = Document()
        document.add_paragraph("   ")
        document.add_paragraph("สัญญาจ้างแรงงาน", style="Heading 1")
        document.add_paragraph("รายละเอียดสัญญา")

        blocks = parse_docx(make_docx(document))

        self.assertEqual(len(blocks), 2)
        self.assertEqual(blocks[0].index, 0)
        self.assertEqual(blocks[0].type, "heading")
        self.assertEqual(blocks[0].style, "Heading 1")
        self.assertEqual(blocks[1].index, 1)
        self.assertEqual(blocks[1].type, "paragraph")


if __name__ == "__main__":
    unittest.main()
